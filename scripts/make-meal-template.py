#!/usr/bin/env python3
"""
Generates HouseholdApp/Resources/MealPlanTemplate.xlsx — the spreadsheet a
household member fills out (Excel, Google Sheets, or Numbers) and then
imports from the Meals tab.

Re-run after changing columns:  python3 scripts/make-meal-template.py
Keep this in sync with MealPlanImporter.swift (column aliases) and the
README's "Plan from a spreadsheet" section.

Error-proofing built into the file itself:
  • Every sheet is protected: header rows and the Instructions sheet are
    locked; only the data cells can be edited.
  • Workbook structure is locked so sheets can't be renamed or deleted.
  • Dropdowns for Meal type and Packing section; dates validated as dates.
  • Header cells carry comments explaining required vs optional.
The importer is also tolerant (case-insensitive headers, many date formats,
blank rows ignored) so mild deviations still work.
"""

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.protection import SheetProtection
import os

OUT = os.path.join(os.path.dirname(__file__), "..", "HouseholdApp", "Resources", "MealPlanTemplate.xlsx")
DATA_ROWS = 200          # how many editable rows to prepare per sheet
PROTECT_PASSWORD = None  # no password: protection is a guardrail, not a lock

MEAL_TYPES = ["Breakfast", "Brunch", "Lunch", "Dinner", "Snack", "Dessert"]
SECTIONS   = ["Clothing", "Food", "Toiletries", "Electronics", "Documents", "Other"]

REQ_FILL  = PatternFill("solid", fgColor="FDE9D9")   # soft orange = required
OPT_FILL  = PatternFill("solid", fgColor="EAF1FB")   # soft blue   = optional
HEAD_FONT = Font(bold=True)
THIN      = Side(style="thin", color="BBBBBB")
BORDER    = Border(bottom=THIN)

# (header, required, width, comment)
MEAL_COLS = [
    ("Date",         True,  13, "REQUIRED on the first row of each meal. Type a date (2026-09-14, Sep 14, 9/14/2026) or use a date cell.\nLeave BLANK on extra ingredient rows for the same meal."),
    ("Meal",         True,  12, "REQUIRED on the first row of each meal. Pick from the dropdown.\nLeave BLANK on extra ingredient rows."),
    ("Meal Name",    False, 24, "Optional dish name, e.g. \"Spaghetti Bolognese\". Blank = just the meal type."),
    ("Cook",         False, 13, "Optional. A household member's name as it appears in the app. Blank = everyone."),
    ("Ingredient",   False, 24, "Optional. ONE ingredient per row. Need more? Use the rows below and leave Date/Meal blank — they belong to the meal above."),
    ("Need to buy?", False, 13, "Pick Yes if this ingredient has to be bought. Yes = a shopping item is created and linked to the meal. Blank or No = you already have it."),
    ("Trip Name",    False, 18, "Optional. Ties the meal to a trip; created for you if it doesn't exist (add dates on the Trips sheet, or it spans the meals' dates). All ingredients go on the trip's packing list under Food."),
    ("Recipe Link",  False, 28, "Optional. Web link to the recipe (https://…)."),
    ("Instructions", False, 36, "Optional. Cooking steps. Line breaks are fine (Alt+Enter in Excel, Ctrl+Enter in Sheets)."),
    ("Notes",        False, 26, "Optional."),
]

# Grey example rows shown at the top of the Meals sheet. The importer skips
# any meal whose name starts with "Example" (and its ingredient rows), so
# leaving them in place is harmless — but the sheet tells people to overwrite.
MEAL_EXAMPLES = [
    ("2026-09-14", "Dinner",    "Example: Spaghetti Bolognese", "Jordan", "ground beef", "Yes", "",                "https://example.com/bolognese", "Brown the beef, add sauce, simmer 20 min.", ""),
    ("",           "",          "",                             "",       "pasta",       "No",  "",                "",                              "",                                         ""),
    ("",           "",          "",                             "",       "garlic",      "",    "",                "",                              "",                                         ""),
    ("2026-09-20", "Breakfast", "Example: Pancakes",            "",       "maple syrup", "Yes", "Cottage weekend", "",                              "",                                         "Bring the good syrup"),
    ("",           "",          "",                             "",       "flour",       "",    "",                "",                              "",                                         ""),
]
PACK_EXAMPLES = [("Cottage weekend", "Example: Rain jacket", "Clothing"), ("Cottage weekend", "Example: Sunscreen", "Toiletries")]
TRIP_EXAMPLES = [("Example: Cottage weekend", "2026-09-19", "2026-09-21", "")]

TRIP_COLS = [
    ("Trip Name",  True,  24, "REQUIRED. Must match the Trip Name used on the Meals and Packing sheets (not case-sensitive)."),
    ("Start Date", True,  14, "REQUIRED. First day of the trip."),
    ("End Date",   True,  14, "REQUIRED. Last day of the trip. Same as Start Date for a single day."),
    ("Notes",      False, 36, "Optional."),
]

PACK_COLS = [
    ("Trip Name", True,  24, "REQUIRED. The trip this item belongs to. Must exist in the app, on the Trips sheet, or be named on a meal."),
    ("Item",      True,  30, "REQUIRED. What to pack, e.g. \"Rain jacket\"."),
    ("Section",   False, 16, "Optional. Pick from the dropdown. Blank = Other."),
]


def write_sheet(ws, cols, freeze="A2"):
    """Header row + formatting + protection with editable data cells."""
    for idx, (title, required, width, comment) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=idx, value=f"{title} *" if required else title)
        cell.font = HEAD_FONT
        cell.fill = REQ_FILL if required else OPT_FILL
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.comment = Comment(comment, "HouseholdApp")
        cell.comment.width = 320
        cell.comment.height = 110
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = freeze

    # Unlock every data cell; the header row stays locked by default.
    for r in range(2, DATA_ROWS + 2):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=r, column=c)
            cell.protection = Protection(locked=False)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.protection = SheetProtection(
        sheet=True, formatCells=False, formatColumns=False, formatRows=False,
        insertRows=False, deleteRows=False, sort=False, autoFilter=False,
        selectLockedCells=False, selectUnlockedCells=False,
    )
    if PROTECT_PASSWORD:
        ws.protection.password = PROTECT_PASSWORD


EXAMPLE_FONT = Font(italic=True, color="9A9A9A")


def write_examples(ws, rows):
    """Grey, italic sample rows right under the header. Cells stay editable."""
    for r, values in enumerate(rows, start=2):
        for c, v in enumerate(values, start=1):
            if v == "":
                continue
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = EXAMPLE_FONT


def add_yes_no(ws, col_letter, title="Need to buy?"):
    """Yes/No dropdown that highlights Yes — the portable stand-in for a
    checkbox (real Excel checkboxes don't survive Google Sheets or Numbers)."""
    from openpyxl.formatting.rule import CellIsRule
    add_list_validation(ws, col_letter, ["Yes", "No"], title)
    rng = f"{col_letter}2:{col_letter}{DATA_ROWS + 1}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Yes"'],
        fill=PatternFill("solid", fgColor="FFD9A8"), font=Font(bold=True, color="8A4B00")))
    for r in range(2, DATA_ROWS + 2):
        ws[f"{col_letter}{r}"].alignment = Alignment(horizontal="center", vertical="top")


def add_list_validation(ws, col_letter, options, title):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.error = f"Pick a {title} from the dropdown."
    dv.errorTitle = f"Invalid {title}"
    dv.prompt = f"Choose a {title}"
    dv.promptTitle = title
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{DATA_ROWS + 1}")


def add_date_validation(ws, col_letter):
    dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True)
    dv.error = "Enter a date, e.g. 2026-09-14."
    dv.errorTitle = "Invalid date"
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{DATA_ROWS + 1}")
    for r in range(2, DATA_ROWS + 2):
        ws[f"{col_letter}{r}"].number_format = "yyyy-mm-dd"


def instructions_sheet(ws):
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 90
    rows = [
        ("HouseholdApp — Meal Plan Template", None),
        ("", None),
        ("How to use", "1. Fill in the Meals sheet. Grey rows are examples — type over them or delete them.\n"
                       "   One ingredient per row: put the Date and Meal on the first row, then list more ingredients on the rows below with Date/Meal left blank.\n"
                       "   Set \"Need to buy?\" to Yes on anything you have to shop for — it becomes a shopping item linked to the meal.\n"
                       "   Trips and Packing sheets are optional.\n"
                       "2. Save the file as .xlsx (Google Sheets: File → Download → Microsoft Excel).\n"
                       "3. In the app: Meals tab → spreadsheet button → Import. You'll see a preview and can fix anything before it's added."),
        ("", None),
        ("Colour key", "Orange header = required.  Blue header = optional.  Hover a header for details."),
        ("Dates", "Type them any common way: 2026-09-14, Sep 14, 9/14/2026, or use the cell as a real date. Past dates are allowed but not recommended."),
        ("Ingredients", "One per row. Rows with a blank Date and Meal belong to the meal above. (Comma-separated lists in one cell still work if you prefer.)"),
        ("Need to buy?", "Yes = create a shopping item. Anything else = you already have it."),
        ("Names", "Cook must match a household member's name in the app (not case-sensitive). Blank = everyone."),
        ("Trips", "Naming a trip on a meal links it. If the trip isn't in the app or on the Trips sheet, it's created spanning the dates of its meals."),
        ("Duplicates", "Meals that already exist (same day, type and name) and shopping/packing items already on a list are skipped, so re-importing the same file is safe."),
        ("Protection", "Headers and sheet names are locked so the import always works. Only the white cells are editable. Don't unprotect unless you know what you're changing."),
        ("", None),
        ("EXAMPLE — Meals", None),
        ("Date | Meal | Meal Name | Cook | Ingredient | Need to buy? | Trip Name",
         "2026-09-14 | Dinner | Spaghetti Bolognese | Jordan | ground beef | Yes |\n"
         "           |        |                     |        | pasta       | No  |\n"
         "           |        |                     |        | garlic      |     |\n"
         "2026-09-20 | Breakfast | Pancakes |  | maple syrup | Yes | Cottage weekend\n"
         "           |           |          |  | flour       |     |"),
        ("", None),
        ("EXAMPLE — Trips", None),
        ("Trip Name | Start Date | End Date", "Cottage weekend | 2026-09-19 | 2026-09-21"),
        ("", None),
        ("EXAMPLE — Packing", None),
        ("Trip Name | Item | Section", "Cottage weekend | Rain jacket | Clothing\nCottage weekend | Sunscreen | Toiletries"),
    ]
    for r, (a, b) in enumerate(rows, start=1):
        ca = ws.cell(row=r, column=1, value=a)
        ca.alignment = Alignment(vertical="top", wrap_text=True)
        if r == 1:
            ca.font = Font(bold=True, size=14)
        elif a and not b:
            ca.font = Font(bold=True)
        elif a:
            ca.font = Font(bold=True, color="444444")
        if b:
            cb = ws.cell(row=r, column=2, value=b)
            cb.alignment = Alignment(vertical="top", wrap_text=True)
    ws.protection = SheetProtection(sheet=True)


def main():
    wb = Workbook()
    ws_i = wb.active
    ws_i.title = "Instructions"
    instructions_sheet(ws_i)

    ws_m = wb.create_sheet("Meals")
    write_sheet(ws_m, MEAL_COLS)
    write_examples(ws_m, MEAL_EXAMPLES)
    add_date_validation(ws_m, "A")
    add_list_validation(ws_m, "B", MEAL_TYPES, "Meal")
    add_yes_no(ws_m, "F")

    ws_t = wb.create_sheet("Trips")
    write_sheet(ws_t, TRIP_COLS)
    write_examples(ws_t, TRIP_EXAMPLES)
    add_date_validation(ws_t, "B")
    add_date_validation(ws_t, "C")

    ws_p = wb.create_sheet("Packing")
    write_sheet(ws_p, PACK_COLS)
    write_examples(ws_p, PACK_EXAMPLES)
    add_list_validation(ws_p, "C", SECTIONS, "Section")

    # Open on Meals; lock sheet structure (no rename/delete/add).
    wb.active = 1
    wb.security = wb.security or None
    from openpyxl.workbook.protection import WorkbookProtection
    wb.security = WorkbookProtection(workbookPassword=None, lockStructure=True)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("wrote", os.path.normpath(OUT))


if __name__ == "__main__":
    main()
